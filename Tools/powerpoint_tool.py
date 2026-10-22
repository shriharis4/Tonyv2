"""
Basic PowerPoint slide creation and editing via python-pptx.

Closes the coverage gap where only a one-way `.pptx -> .pdf` converter
(`word_to_pdf.ppt_to_pdf`) existed, with no way to author a presentation.
"""

from __future__ import annotations

import asyncio
import os
from typing import List, Literal, Optional

from livekit.agents import function_tool

from Tools.tool_result import ok, fail

DEFAULT_DOC_DIR = os.path.join(os.path.expanduser("~"), "Documents")


def _resolve_pptx_path(path: str) -> str:
    path = os.path.expandvars(os.path.expanduser((path or "presentation.pptx").strip()))
    if not os.path.isabs(path):
        path = os.path.join(DEFAULT_DOC_DIR, path)
    if not path.lower().endswith(".pptx"):
        path += ".pptx"
    return path


def _build_presentation(path: str, title: str, slides_content: List[str]) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    bullet_slide_layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title or "Presentation"

    for content in slides_content:
        slide = prs.slides.add_slide(bullet_slide_layout)
        lines = [l.strip() for l in content.split("\n") if l.strip()]
        if not lines:
            continue
        slide.shapes.title.text = lines[0]
        if len(lines) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.text = lines[1]
            for extra in lines[2:]:
                p = tf.add_paragraph()
                p.text = extra

    os.makedirs(os.path.dirname(path), exist_ok=True)
    prs.save(path)


@function_tool()
async def create_presentation(path: str, title: str, slides_content: str) -> str:
    """
    Create a new PowerPoint presentation with a title slide and one slide
    per section of content.

    Args:
        path: File name or path for the .pptx file. Relative names are
            created under the user's Documents folder.
        title: Title shown on the opening slide.
        slides_content: Slide bodies separated by "---". Within each slide,
            the first line is the slide title and subsequent lines become
            bullet points.

    Returns:
        Structured success/failure string with the final file path.
    """
    if not slides_content or not slides_content.strip():
        return fail("No slide content provided.")

    try:
        final_path = _resolve_pptx_path(path)
        sections = [s for s in slides_content.split("---")]
        loop = asyncio.get_running_loop()
        await loop.run_in_executor(None, _build_presentation, final_path, title, sections)
        return ok(f"Presentation created at {final_path}", detail=f"slides={len(sections) + 1}")
    except ImportError:
        return fail("python-pptx is required for PowerPoint authoring.", detail="pip install python-pptx")
    except Exception as e:
        return fail("Failed to create presentation", detail=str(e))


def _edit_presentation(path: str, operation: str, detail: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)

    if operation == "add_slide":
        title, _, body = detail.partition("=>")
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title.strip()
        if body.strip():
            slide.placeholders[1].text_frame.text = body.strip()
        prs.save(path)
        return f"Added slide '{title.strip()}'"

    elif operation == "delete_slide":
        index = int(detail.strip())
        if index < 0 or index >= len(prs.slides):
            raise ValueError(f"Slide index {index} out of range (0-{len(prs.slides) - 1})")
        xml_slides = prs.slides._sldIdLst
        slides = list(xml_slides)
        xml_slides.remove(slides[index])
        prs.save(path)
        return f"Deleted slide {index}"

    elif operation == "edit_title":
        index_str, _, new_title = detail.partition("=>")
        index = int(index_str.strip())
        if index < 0 or index >= len(prs.slides):
            raise ValueError(f"Slide index {index} out of range (0-{len(prs.slides) - 1})")
        slide = prs.slides[index]
        if slide.shapes.title:
            slide.shapes.title.text = new_title.strip()
        prs.save(path)
        return f"Updated title of slide {index}"

    else:
        raise ValueError(f"Unknown operation '{operation}'")


@function_tool()
async def edit_presentation(
    path: str,
    operation: Literal["add_slide", "delete_slide", "edit_title"],
    detail: str,
) -> str:
    """
    Edit an existing PowerPoint presentation.

    Args:
        path: File name or path to an existing .pptx file.
        operation: "add_slide" (detail "TITLE=>BODY"), "delete_slide"
            (detail is the 0-indexed slide number), or "edit_title"
            (detail "INDEX=>NEW_TITLE").
        detail: Operation-specific payload, see `operation` above.

    Returns:
        Structured success/failure string.
    """
    final_path = os.path.expandvars(os.path.expanduser((path or "").strip()))
    if not os.path.isfile(final_path):
        return fail(f"Presentation not found: {final_path}")

    try:
        loop = asyncio.get_running_loop()
        msg = await loop.run_in_executor(None, _edit_presentation, final_path, operation, detail)
        return ok(msg, detail=f"file={final_path}")
    except ValueError as e:
        return fail(str(e))
    except ImportError:
        return fail("python-pptx is required to edit presentations.", detail="pip install python-pptx")
    except Exception as e:
        return fail("Failed to edit presentation", detail=str(e))
